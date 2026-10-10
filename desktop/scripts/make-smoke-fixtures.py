"""Create small real documents; these fixtures are never shipped to users."""
import argparse
import io
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


def create_fixtures(destination: Path, font_path: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    shared = "本项目采用分阶段实施方案，建立质量检查机制，确保工程按计划完成。"
    reference = "施工组织方案覆盖质量安全和进度，所有供应商必须遵守统一验收要求。"
    picture = io.BytesIO()
    Image.new("RGB", (80, 80), "#547c65").save(picture, format="PNG")
    for name in ("甲", "乙"):
        document = Document()
        document.add_paragraph(shared)
        document.add_paragraph(reference)
        document.add_paragraph(f"{name}标书另行安排项目负责人和现场施工管理。")
        document.add_picture(io.BytesIO(picture.getvalue()))
        document.save(destination / f"{name}.docx")
    tender = Document()
    tender.add_paragraph(reference)
    tender.save(destination / "招标文件.docx")
    text_pdf = canvas.Canvas(str(destination / "文字.pdf"))
    text_pdf.drawString(60, 730, "The project includes quality inspection and phased construction.")
    text_pdf.save()
    scan = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(scan)
    font = ImageFont.truetype(str(font_path), 40)
    for index, text in enumerate([
        "标书质量检查和施工组织方案。",
        "项目采用分阶段实施，确保工程按计划完成。",
        "QUALITY INSPECTION FOR CONSTRUCTION PROJECT",
    ]):
        draw.text((70, 150 + index * 100), text, font=font, fill="black")
    scan_pdf = canvas.Canvas(str(destination / "扫描.pdf"))
    scan_pdf.drawImage(ImageReader(scan), 0, 0, width=595, height=842)
    scan_pdf.save()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--font", type=Path, required=True)
    args = parser.parse_args()
    create_fixtures(args.output, args.font)
