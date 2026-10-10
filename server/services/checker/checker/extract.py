import base64
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import BinaryIO, Literal

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, UnidentifiedImageError

from checker.schemas import ImageInfo, TextBlock

MAX_UNCOMPRESSED = 2 * 1024 * 1024 * 1024
MAX_MEMBERS = 10000
MAX_PREVIEW_BYTES = 20 * 1024 * 1024


@dataclass(frozen=True)
class ParsedDocument:
    id: str
    name: str
    blocks: list[TextBlock]
    images: list[ImageInfo]
    format: Literal["docx", "pdf"] = "docx"
    pages: int | None = None
    ocr_pages: int = 0
    warnings: tuple[str, ...] = ()


def _preview_and_hash(data: bytes) -> tuple[str | None, int | None, tuple[int, int, int] | None, float | None, str | None]:
    try:
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 40_000_000:
                return None, None, None, None, None
            image.load()
            rgb = image.convert("RGB")
            pixel_hash = None
            if image.width * image.height <= 10_000_000:
                pixel_hash = hashlib.sha256(f"{image.width}x{image.height}:".encode() + rgb.tobytes()).hexdigest()
            gray = image.convert("L").resize((9, 8))
            pixels = list(gray.get_flattened_data())
            average = sum(pixels) / len(pixels)
            variance = sum((pixel - average) ** 2 for pixel in pixels) / len(pixels)
            mean = rgb.resize((1, 1)).getpixel((0, 0))
            bits = sum(
                1 << (row * 8 + col)
                for row in range(8)
                for col in range(8)
                if pixels[row * 9 + col] > pixels[row * 9 + col + 1]
            )
            image.thumbnail((320, 200))
            if image.mode not in ("RGB", "L"):
                image = image.convert("RGB")
            out = io.BytesIO()
            image.save(out, format="JPEG", quality=72)
            return "data:image/jpeg;base64," + base64.b64encode(out.getvalue()).decode(), bits, mean, variance, pixel_hash
    except (Image.DecompressionBombError, UnidentifiedImageError, OSError, ValueError):
        return None, None, None, None, None


def image_info(data: bytes, name: str) -> ImageInfo:
    if len(data) > MAX_PREVIEW_BYTES:
        return ImageInfo(sha256=hashlib.sha256(data).hexdigest(), name=name)
    preview, perceptual_hash, mean_rgb, visual_variance, pixel_sha256 = _preview_and_hash(data)
    return ImageInfo(
        sha256=hashlib.sha256(data).hexdigest(),
        name=name,
        preview=preview,
        perceptual_hash=perceptual_hash,
        mean_rgb=mean_rgb,
        visual_variance=visual_variance,
        pixel_sha256=pixel_sha256,
    )


def _as_stream(data: bytes | BinaryIO) -> BinaryIO:
    stream = io.BytesIO(data) if isinstance(data, bytes) else data
    stream.seek(0)
    return stream


def _images(data: bytes | BinaryIO) -> list[ImageInfo]:
    images: list[ImageInfo] = []
    with zipfile.ZipFile(_as_stream(data)) as archive:
        members = archive.infolist()
        if len(members) > MAX_MEMBERS or sum(m.file_size for m in members) > MAX_UNCOMPRESSED:
            raise ValueError("DOCX 解压后内容过大或文件项过多")
        for member in members:
            if not member.filename.startswith("word/media/") or member.is_dir():
                continue
            if member.file_size > MAX_PREVIEW_BYTES:
                digest = hashlib.sha256()
                with archive.open(member) as media:
                    for chunk in iter(lambda: media.read(1024 * 1024), b""):
                        digest.update(chunk)
                images.append(ImageInfo(sha256=digest.hexdigest(), name=PurePosixPath(member.filename).name))
                continue
            raw = archive.read(member)
            images.append(image_info(raw, PurePosixPath(member.filename).name))
    return images


def _append_text(blocks: list[TextBlock], text: str, location: str) -> None:
    text = re.sub(r"\s+", " ", text).strip()
    if text:
        blocks.append(TextBlock(text=text, location=location))


def _table_blocks(table: Table, location: str, blocks: list[TextBlock]) -> None:
    seen_cells: set[int] = set()
    for row_number, row in enumerate(table.rows, start=1):
        for col_number, cell in enumerate(row.cells, start=1):
            cell_identity = id(cell._tc)
            if cell_identity in seen_cells:
                continue
            seen_cells.add(cell_identity)
            cell_location = f"{location} · 行 {row_number} 列 {col_number}"
            for child in cell.iter_inner_content():
                if isinstance(child, Paragraph):
                    _append_text(blocks, child.text, cell_location)
                elif isinstance(child, Table):
                    _table_blocks(child, cell_location, blocks)


def parse_docx(data: bytes | BinaryIO, name: str, document_id: str) -> ParsedDocument:
    try:
        stream = _as_stream(data)
        images = _images(stream)
        stream.seek(0)
        document = Document(stream)
    except ValueError:
        raise
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("文件不是有效的 .docx 文档") from exc
    blocks: list[TextBlock] = []
    paragraph_number = 0
    table_number = 0
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            paragraph_number += 1
            _append_text(blocks, item.text, f"正文 · 段落 {paragraph_number}")
        elif isinstance(item, Table):
            table_number += 1
            _table_blocks(item, f"表格 {table_number}", blocks)
    for section_number, section in enumerate(document.sections, start=1):
        for label, part in (("页眉", section.header), ("页脚", section.footer)):
            for paragraph in part.paragraphs:
                _append_text(blocks, paragraph.text, f"第 {section_number} 节{label}")
    return ParsedDocument(id=document_id, name=name, blocks=blocks, images=images)
