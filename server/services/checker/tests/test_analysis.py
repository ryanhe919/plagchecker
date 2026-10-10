import io
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from checker.api import app
from checker.compare import analyze
from checker.extract import ParsedDocument
from checker.schemas import ImageInfo, Settings, TextBlock


def make_docx(paragraphs: list[str], image: bytes | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "施工组织方案覆盖质量、安全和进度三个方面。"
    if image:
        document.add_picture(io.BytesIO(image))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def make_image(color: str) -> bytes:
    image = Image.new("RGB", (80, 80), color)
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def make_pdf(pages: list[str], image: bytes | None = None) -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    stream = io.BytesIO()
    document = canvas.Canvas(stream)
    for text in pages:
        document.setFont("STSong-Light", 14)
        document.drawString(60, 730, text)
        if image:
            document.drawImage(ImageReader(io.BytesIO(image)), 60, 590, width=90, height=90)
        document.showPage()
    document.save()
    return stream.getvalue()


def make_scanned_pdf(image: bytes) -> bytes:
    stream = io.BytesIO()
    document = canvas.Canvas(stream)
    document.drawImage(ImageReader(io.BytesIO(image)), 0, 0, width=595, height=842)
    document.save()
    return stream.getvalue()


def test_analysis_finds_text_and_byte_identical_image() -> None:
    shared = "本项目采用分阶段实施方案，建立质量检查机制，确保工程按计划完成。"
    image = make_image("#547c65")
    first = make_docx([shared, "第一份标书独有的施工说明。"], image)
    second = make_docx([shared, "第二份标书独有的成本说明。"], image)
    third = make_docx(["第三份标书提出完全不同的技术路线和验收方式。"], make_image("#873e42"))
    response = TestClient(app).post(
        "/api/v1/checker/analyses",
        files=[("files", ("甲.docx", first)), ("files", ("乙.docx", second)), ("files", ("丙.docx", third))],
        data={"settings": "{}"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert len(body["pairs"]) == 3
    pair = next(item for item in body["pairs"] if {item["a_name"], item["b_name"]} == {"甲.docx", "乙.docx"})
    assert pair["similarity"] > 0.4
    assert pair["exact_image_count"] == 1
    assert any(shared in passage["a_text"] for passage in pair["passages"])
    assert all(passage["a_highlights"] and passage["b_highlights"] for passage in pair["passages"])
    for passage in pair["passages"]:
        a_doc = next(item for item in body["documents"] if item["id"] == pair["a_id"])
        b_doc = next(item for item in body["documents"] if item["id"] == pair["b_id"])
        assert a_doc["content"][passage["a_block_index"]]["text"][passage["a_start"] :].startswith(passage["a_text"])
        assert b_doc["content"][passage["b_block_index"]]["text"][passage["b_start"] :].startswith(passage["b_text"])
    assert all(not item["images"] for item in body["pairs"] if "丙.docx" in (item["a_name"], item["b_name"]))


def test_common_template_can_be_excluded() -> None:
    template = "施工组织方案覆盖质量、安全和进度三个方面。"
    documents = [make_docx([template, name]) for name in ("甲方要求另行记录施工条件。", "乙方要求另行记录项目成本。", "丙方要求另行记录验收标准。")]
    response = TestClient(app).post(
        "/api/v1/checker/analyses",
        files=[("files", (f"{index}.docx", item)) for index, item in enumerate(documents)],
        data={"settings": '{"ignore_common_across_all":true}'},
    )
    assert response.status_code == 200, response.text
    assert all(pair["similarity"] == 0 for pair in response.json()["pairs"])


def test_invalid_file_is_rejected() -> None:
    response = TestClient(app).post(
        "/api/v1/checker/analyses",
        files=[("files", ("甲.docx", b"not a zip")), ("files", ("乙.docx", b"not a zip"))],
        data={"settings": "{}"},
    )
    assert response.status_code == 422


def test_similar_but_not_identical_passage_respects_threshold() -> None:
    first = "我们计划建立施工现场安全检查与阶段质量复核制度，确保每个关键节点均有责任人签字确认。"
    second = "我们拟建立施工现场安全检查及阶段质量复核制度，确保各关键节点均由负责人签字确认。"
    files = [("files", ("甲.docx", make_docx([first]))), ("files", ("乙.docx", make_docx([second])))]
    client = TestClient(app)
    normal = client.post("/api/v1/checker/analyses", files=files, data={"settings": "{}"})
    strict = client.post("/api/v1/checker/analyses", files=files, data={"settings": '{"similarity_threshold":1}'})
    assert normal.status_code == strict.status_code == 200
    assert any(not item["exact"] for item in normal.json()["pairs"][0]["passages"])
    assert all(item["exact"] for item in strict.json()["pairs"][0]["passages"])


def test_passage_offsets_point_into_complete_paragraph() -> None:
    shared = "工程实施时采用分阶段质量验收机制，并在每个节点保留书面签字记录。"
    first = make_docx(["甲方的前置说明。" + shared])
    second = make_docx(["乙方的不同前置说明。" + shared])
    response = TestClient(app).post(
        "/api/v1/checker/analyses",
        files=[("files", ("甲.docx", first)), ("files", ("乙.docx", second))],
        data={"settings": "{}"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    passage = next(item for item in body["pairs"][0]["passages"] if shared in item["a_text"])
    assert passage["a_start"] > 0 and passage["b_start"] > 0
    assert body["documents"][0]["content"][passage["a_block_index"]]["text"][passage["a_start"] :].startswith(shared)
    assert body["documents"][1]["content"][passage["b_block_index"]]["text"][passage["b_start"] :].startswith(shared)


def test_pdf_and_docx_can_be_compared_with_page_locations_and_images() -> None:
    shared = "本项目采用分阶段施工计划，建立质量检查制度并对重要节点留存记录。"
    image = make_image("#547c65")
    docx = make_docx([shared], image)
    pdf = make_pdf(["第一页是本文件的独立说明。", shared], image)
    response = TestClient(app).post(
        "/api/v1/checker/analyses",
        files=[("files", ("甲.docx", docx)), ("files", ("乙.pdf", pdf))],
        data={"settings": '{"pdf_ocr_mode":"off"}'},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["documents"][1]["format"] == "pdf"
    assert body["documents"][1]["pages"] == 2
    pair = body["pairs"][0]
    assert any(item["b_location"] == "PDF · 第 2 页" for item in pair["passages"])
    assert pair["exact_image_count"] == 1
    assert pair["images"][0]["match_type"] in {"byte_exact", "pixel_exact"}


def test_invalid_pdf_is_rejected() -> None:
    valid = make_docx(["这是一个有效的 DOCX 标书正文，供错误 PDF 测试使用。"])
    response = TestClient(app).post(
        "/api/v1/checker/analyses",
        files=[("files", ("甲.docx", valid)), ("files", ("乙.pdf", b"not a pdf"))],
        data={"settings": "{}"},
    )
    assert response.status_code == 422


def test_pdf_auto_ocr_and_off_mode_are_reported() -> None:
    shared = "本项目建立施工质量复核流程，并对重要节点保存现场记录。"
    scan = make_scanned_pdf(make_image("#547c65"))
    docx = make_docx([shared])
    files = [("files", ("甲.docx", docx)), ("files", ("扫描.pdf", scan))]
    with patch("checker.extract_pdf._ocr_text", return_value=shared) as ocr:
        auto = TestClient(app).post("/api/v1/checker/analyses", files=files, data={"settings": "{}"})
    off = TestClient(app).post("/api/v1/checker/analyses", files=files, data={"settings": '{"pdf_ocr_mode":"off"}'})
    assert auto.status_code == off.status_code == 200
    assert ocr.call_count == 1
    assert auto.json()["documents"][1]["ocr_pages"] == 1
    assert auto.json()["pairs"][0]["passages"]
    assert off.json()["documents"][1]["ocr_pages"] == 0
    assert off.json()["documents"][1]["warnings"]


def test_pixel_identical_images_with_different_encodings_are_separate_from_byte_matches() -> None:
    picture = Image.new("RGB", (96, 64))
    for x in range(96):
        for y in range(64):
            picture.putpixel((x, y), (x * 2, y * 3, (x + y) % 256))
    first_image = io.BytesIO()
    second_image = io.BytesIO()
    picture.save(first_image, format="PNG", compress_level=1)
    picture.save(second_image, format="PNG", compress_level=9)
    assert first_image.getvalue() != second_image.getvalue()
    first = make_docx(["第一份标书采用分阶段质量管理的方案。"], first_image.getvalue())
    second = make_docx(["第二份标书采用独立的管理方案。"], second_image.getvalue())
    response = TestClient(app).post(
        "/api/v1/checker/analyses",
        files=[("files", ("甲.docx", first)), ("files", ("乙.docx", second))],
        data={"settings": "{}"},
    )
    assert response.status_code == 200, response.text
    pair = response.json()["pairs"][0]
    assert pair["exact_image_count"] == 1
    assert pair["images"][0]["match_type"] == "pixel_exact"


def test_repeated_exact_passages_map_to_distinct_source_blocks() -> None:
    phrase = "施工质量检查记录需要逐项核实并由负责人签字确认。"
    a = ParsedDocument("1", "甲.pdf", [TextBlock(text=phrase, location=f"第 {index} 页") for index in range(2)], [])
    b = ParsedDocument("2", "乙.pdf", [TextBlock(text=phrase, location=f"第 {index} 页") for index in range(3)], [])
    pair = analyze([a, b], Settings()).pairs[0]
    assert len(pair.passages) == 2
    assert {passage.b_block_index for passage in pair.passages} == {0, 1}
    assert pair.a_coverage == 1
    assert 0.6 < pair.b_coverage < 0.7


def test_image_count_includes_matches_beyond_display_limit() -> None:
    images_a = [ImageInfo(sha256=f"hash-{index}", name=f"a-{index}.png") for index in range(120)]
    images_b = [ImageInfo(sha256=f"hash-{index}", name=f"b-{index}.png") for index in range(120)]
    a = ParsedDocument("1", "甲.pdf", [], images_a)
    b = ParsedDocument("2", "乙.pdf", [], images_b)
    pair = analyze([a, b], Settings()).pairs[0]
    assert pair.exact_image_count == 120
    assert len(pair.images) == 100


def test_all_text_matches_are_returned_beyond_previous_200_limit() -> None:
    blocks = [
        TextBlock(text=f"第{index}项施工过程质量检查要求由项目负责人签字确认。", location=f"第 {index} 页")
        for index in range(250)
    ]
    first = ParsedDocument("1", "甲.pdf", blocks, [])
    second = ParsedDocument("2", "乙.pdf", blocks, [])
    pair = analyze([first, second], Settings()).pairs[0]
    assert len(pair.passages) == 250
    assert any(passage.a_block_index == 249 for passage in pair.passages)


def test_tender_reference_excludes_only_shared_reference_text() -> None:
    tender = "招标文件要求建立设备安装质量检查流程并保留验收记录"
    unique = "另行采用企业自主研发的智能监测与异常预警方法"
    text = tender + unique + "。"
    first = ParsedDocument("1", "甲.docx", [TextBlock(text=text, location="正文")], [])
    second = ParsedDocument("2", "乙.docx", [TextBlock(text=text, location="正文")], [])
    reference = ParsedDocument("r", "招标.pdf", [TextBlock(text=tender, location="第 1 页")], [])
    pair = analyze([first, second], Settings(), reference=reference).pairs[0]
    assert pair.passages
    assert pair.passages[0].matched_chars == len(unique)
    assert pair.passages[0].a_highlights[0][0] >= len(tender)
    assert 0.99 < pair.similarity <= 1
    result = analyze([first, second], Settings(), reference=reference)
    assert result.reference is not None and result.reference.name == "招标.pdf"
    assert result.documents[0].excluded_reference_chars >= len(tender)
    assert result.documents[0].content[0].reference_excluded_ranges == [(0, len(tender))]


def test_tender_reference_images_are_not_reported_as_bid_duplication() -> None:
    tender_image = ImageInfo(sha256="tender-image", name="tender.png")
    unique_image = ImageInfo(sha256="unique-image", name="unique.png")
    reference = ParsedDocument("r", "招标.pdf", [], [tender_image])
    first = ParsedDocument("1", "甲.docx", [], [tender_image, unique_image])
    second = ParsedDocument("2", "乙.docx", [], [tender_image, unique_image])
    result = analyze([first, second], Settings(), reference=reference)
    assert result.pairs[0].exact_image_count == 1
    assert result.pairs[0].images[0].sha256 == "unique-image"
    assert result.documents[0].excluded_reference_images == 1
